"""Fabrique le modèle Excel à remplir avec les forfaits.

    python outils/modele_excel.py            ->  prive/Forfaits-a-remplir.xlsx   (ignoré par git)
    python outils/modele_excel.py autre.xlsx ->  autre.xlsx

Refuse d'écraser un fichier qui existe déjà : celui de `prive/` est celui qu'on remplit à la main.

Le tableau se remplit dans Excel, s'enregistre en « CSV UTF-8 », puis s'importe dans l'outil
(Menu -> Importer). Les titres de colonnes sont ceux de l'export CSV de l'outil : ne pas les changer
sans changer `COLONNES_TABLEAU` dans index.html.
"""

import pathlib
import sys

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

RACINE = pathlib.Path(__file__).resolve().parent.parent
SORTIE = RACINE / "prive" / "Forfaits-a-remplir.xlsx"

DERNIERE_LIGNE = 1000  # nombre de forfaits que le modèle sait contrôler
DERNIERE_LIGNE_LISTE = 100  # éléments par famille dans l'onglet Listes
DERNIERE_COLONNE = 26  # colonne Z : jusqu'où une ligne compte comme « remplie » (familles ajoutées à droite comprises)
TVA = 1.2
TOLERANCE_TVA = 0.02  # même tolérance que l'outil : 2 centimes

# Familles de départ : celles de l'outil, listes vides (les éléments de départ de l'outil sont des exemples :
# les vrais se tapent dans l'onglet Listes). L'import retrouve les colonnes par leur titre : l'ordre ici est libre.
# Les familles qui servent rarement sont à droite, pour que celles du quotidien se suivent.
FAMILLES: dict[str, list[str]] = {
    "Taximètre": [],
    "Imprimante": [],
    "Lumineux": [],
    "Boîtier": [],
    "Installation": [],
}
COLONNES_FIXES = ["Référence", "Désignation", "Prix HT", "Prix TTC", "Surnoms"]
PAS_DE_DONNEES = "/"  # en tête de chaque liste : l'outil le lit comme une case vide

ROUGE = "E30613"
NOIR = "000000"
BLANC = "FFFFFF"

MODE_D_EMPLOI = [
    (
        "1. Complète les listes",
        "Onglet « Listes » : un taximètre, boîtier, imprimante… par ligne, sous le titre de sa famille (le « / » en tête veut dire « pas de données » : ne pas l'enlever).",
    ),
    (
        "2. Une ligne par forfait",
        "Onglet « Forfaits » : référence, désignation (le contenu), prix HT et prix TTC tels qu'écrits sur le tarif. "
        "Puis choisis chaque élément dans la liste déroulante. Si le forfait n'en a pas : « / » (pas de données) ou case vide, c'est pareil. "
        "Un forfait sans taximètre (dash cam, lumineux…) n'a besoin que de la référence, de la désignation et des prix. "
        "Boîtier et Installation sont à droite parce qu'ils servent rarement. "
        "Surnoms (facultatif) : les noms que les techniciens donnent au forfait, séparés par des virgules (« le gros pack, pack nuit »), 3 au maximum ; "
        "la recherche les reconnaît sur tous les postes. Un surnom ne désigne qu'une seule chose : un surnom déjà pris est ignoré (annoncé à l'import).",
    ),
    (
        "3. Les couleurs aident",
        "Case rouge = à compléter, ou référence en double. Case orange = le TTC ne correspond pas au HT + 20 %. "
        "Rien d'autre à faire : tu peux laisser une orange si le tarif officiel est ainsi.",
    ),
    (
        "4. Enregistre en CSV",
        "Reste sur l'onglet « Forfaits », puis Fichier > Enregistrer sous > type « CSV UTF-8 (délimité par des virgules) (*.csv) ». "
        "Excel prévient que seul l'onglet actif est enregistré : réponds Oui.",
    ),
    (
        "5. Importe dans l'outil",
        "Menu > « Importer (sauvegarde ou tableau Excel) » > choisis le fichier .csv. L'outil montre ce qui va changer avant de le faire. "
        "Au premier remplissage, choisis « Importer et supprimer les exemples ».",
    ),
    (
        "6. Partage",
        "Ensuite, dans l'outil : Menu > Exporter > Fichier de sauvegarde (.json), à ranger dans le dossier OneDrive des forfaits. "
        "Les autres appareils font Importer > Mettre à jour.",
    ),
    (
        "Un élément ou une famille en plus",
        "Élément nouveau : tape-le directement dans la case (Excel prévient qu'il n'est pas dans la liste : garde-le). "
        "Famille nouvelle (dash cam, TPE…) : écris son titre dans la première case libre à droite de la ligne 1, puis remplis la colonne. "
        "À l'import, l'outil crée la famille et les éléments, et te les montre avant de valider : vérifie l'orthographe à ce moment-là.",
    ),
    (
        "Pour modifier plus tard",
        "Retape la référence d'un forfait existant : il est mis à jour au lieu d'être dupliqué. "
        "Une ligne supprimée du tableau ne supprime rien dans l'outil (supprime dans l'outil).",
    ),
]


def fabriquer() -> Workbook:
    classeur = Workbook()
    forfaits = classeur.active
    forfaits.title = "Forfaits"
    listes = classeur.create_sheet("Listes")
    aide = classeur.create_sheet("Mode d'emploi")

    entete_fond = PatternFill("solid", start_color=ROUGE)
    entete_police = Font(bold=True, color=BLANC, size=12)

    # ----- Listes : une colonne par famille, un élément par ligne
    for col, (famille, elements) in enumerate(FAMILLES.items(), start=1):
        cellule = listes.cell(row=1, column=col, value=famille)
        cellule.fill = PatternFill("solid", start_color=NOIR)
        cellule.font = entete_police
        listes.column_dimensions[get_column_letter(col)].width = 26
        for ligne, element in enumerate([PAS_DE_DONNEES, *elements], start=2):
            listes.cell(row=ligne, column=col, value=element)
    listes.freeze_panes = "A2"

    # ----- Forfaits
    titres = COLONNES_FIXES + list(FAMILLES)
    for col, titre in enumerate(titres, start=1):
        cellule = forfaits.cell(row=1, column=col, value=titre)
        cellule.fill = entete_fond
        cellule.font = entete_police
        cellule.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
    forfaits.row_dimensions[1].height = 24
    forfaits["A1"].comment = Comment(
        "Mode d'emploi : onglet « Mode d'emploi ». Ne change pas les titres de cette ligne.", "Forfaits"
    )
    forfaits.freeze_panes = "A2"
    largeurs = {"A": 22, "B": 46, "C": 13, "D": 13, "E": 30}
    for col in range(1, len(titres) + 1):
        lettre = get_column_letter(col)
        forfaits.column_dimensions[lettre].width = largeurs.get(lettre, 24)

    for ligne in range(2, DERNIERE_LIGNE + 1):
        forfaits.cell(row=ligne, column=1).number_format = "@"  # texte : « 001 » ou « 1E5 » restent tels quels
        forfaits.cell(row=ligne, column=2).alignment = Alignment(wrap_text=True, vertical="top")
        for col in (3, 4):
            forfaits.cell(row=ligne, column=col).number_format = "0.00"

    # Listes déroulantes : un élément de la famille, ou rien. La liste suit l'onglet Listes.
    premiere_famille = len(COLONNES_FIXES) + 1
    for decalage in range(len(FAMILLES)):
        lettre_liste = get_column_letter(decalage + 1)
        formule = f"OFFSET(Listes!${lettre_liste}$2,0,0,MAX(1,COUNTA(Listes!${lettre_liste}$2:${lettre_liste}${DERNIERE_LIGNE_LISTE})),1)"
        validation = DataValidation(
            type="list",
            formula1="=" + formule,
            allow_blank=True,
            showErrorMessage=True,
            errorStyle="warning",  # avertit mais laisse passer : un élément nouveau se tape, l'outil le crée à l'import
            errorTitle="Élément pas dans la liste",
            error="Cet élément n'est pas dans l'onglet « Listes ». L'outil le créera à l'import. Garder cette valeur ?",
        )
        lettre = get_column_letter(premiere_famille + decalage)
        validation.add(f"{lettre}2:{lettre}{DERNIERE_LIGNE}")
        forfaits.add_data_validation(validation)

    # Contrôles en couleur
    rouge = PatternFill("solid", start_color="F8B4BA", end_color="F8B4BA")
    orange = PatternFill("solid", start_color="FFD8A8", end_color="FFD8A8")
    ligne_remplie = f"COUNTA($A2:${get_column_letter(DERNIERE_COLONNE)}2)>0"  # compte aussi les colonnes ajoutées à droite
    for plage, cellule in (("A", "A2"), ("C", "C2"), ("D", "D2")):
        forfaits.conditional_formatting.add(
            f"{plage}2:{plage}{DERNIERE_LIGNE}", FormulaRule(formula=[f'AND({ligne_remplie},{cellule}="")'], fill=rouge)
        )
    forfaits.conditional_formatting.add(
        f"A2:A{DERNIERE_LIGNE}",
        FormulaRule(formula=[f'AND($A2<>"",COUNTIF($A$2:$A${DERNIERE_LIGNE},$A2)>1)'], fill=rouge),
    )
    forfaits.conditional_formatting.add(
        f"D2:D{DERNIERE_LIGNE}",
        FormulaRule(formula=[f"AND(ISNUMBER($C2),ISNUMBER($D2),ABS($D2-$C2*{TVA})>{TOLERANCE_TVA})"], fill=orange),
    )

    # ----- Mode d'emploi
    aide.column_dimensions["A"].width = 28
    aide.column_dimensions["B"].width = 110
    aide["A1"] = "Remplir les forfaits avec Excel"
    aide["A1"].font = Font(bold=True, size=14, color=ROUGE)
    for ligne, (titre, texte) in enumerate(MODE_D_EMPLOI, start=3):
        aide.cell(row=ligne, column=1, value=titre).font = Font(bold=True)
        cellule = aide.cell(row=ligne, column=2, value=texte)
        cellule.alignment = Alignment(wrap_text=True, vertical="top")
        aide.cell(row=ligne, column=1).alignment = Alignment(vertical="top")

    classeur.active = 0  # le CSV enregistré est celui de l'onglet actif : « Forfaits »
    return classeur


def main() -> None:
    sortie = pathlib.Path(sys.argv[1]) if len(sys.argv) > 1 else SORTIE
    if sortie.exists():  # le fichier est souvent déjà rempli à la main : jamais d'écrasement
        sys.exit(f"Refusé : {sortie} existe déjà (peut-être déjà rempli). Donne un autre chemin, ou supprime-le toi-même.")
    sortie.parent.mkdir(parents=True, exist_ok=True)
    fabriquer().save(sortie)
    print(f"Modèle écrit : {sortie}")


if __name__ == "__main__":
    main()
